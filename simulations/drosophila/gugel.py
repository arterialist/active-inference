"""Gugel et al. (2023), Figure 7: primary PN current and firing measurements.

Original workbooks stay unmodified. Cached formula values are independently
checked, not accepted as recalculation. Source columns are experimental cells,
not FlyWire identities. The kernel comparison below does not assign pA to PAULA.
"""
from __future__ import annotations

import argparse
import inspect
import json
import math
from pathlib import Path
import re

import numpy as np

if __package__:
    from .prisco import digest, dump_new, number, verify_file
else:  # Read-only extraction in the bundled document runtime, without MuJoCo.
    from prisco import digest, dump_new, number, verify_file

FILENAME = "elife-85443-fig7-data1-v2.xlsx"
SOURCE = {
    "filename": FILENAME,
    "url": "https://cdn.elifesciences.org/articles/85443/" + FILENAME,
    "article": "https://doi.org/10.7554/eLife.85443",
    "article_version": 2,
    "license": "CC-BY-4.0",
    "attribution": "Zhannetta V Gugel, Elizabeth G Maurais, Elizabeth J Hong (2023)",
    "bytes": 402813,
    "sha256": "a8ae6fcd3bf0d8effab7a0ecbfa88fccf192f134144072282758ca8125bd8c78",
    "hash_origin": "Computed from original publisher download, not a repository-supplied digest",
}


def check_formulas(values: dict, formulas: dict, *, atol: float = 1e-10) -> dict:
    """Evaluate only the published reference +/- literal grammar, without eval.

    Resolve dependencies from constants, detect cycles and reject unsupported
    formulas. Never repair the source or replace cached values in place.
    """
    resolved, visiting = {}, set()

    def resolve(address):
        if address in resolved:
            return resolved[address]
        if address in visiting:
            raise ValueError(f"Formula cycle at {address}")
        if address not in values:
            raise ValueError(f"Missing referenced cell {address}")
        visiting.add(address)
        if address in formulas:
            match = re.fullmatch(r"=([A-Z]+[1-9][0-9]*)([+-])([0-9]+(?:\.[0-9]+)?)",
                                 formulas[address])
            if match is None:
                raise ValueError(f"Unsupported source formula at {address}")
            reference, op, literal = match.groups()
            value = resolve(reference) + (1 if op == "+" else -1) * float(literal)
        else:
            value = number(values[address])
        visiting.remove(address)
        resolved[address] = value
        return value

    residuals = {}
    for address in formulas:
        residuals[address] = number(values[address]) - resolve(address)
        if abs(residuals[address]) > atol:
            raise ValueError(f"Cached formula disagrees at {address}")
    return {"formula_count": len(formulas), "tolerance": atol,
            "maximum_absolute_residual": max(map(abs, residuals.values()), default=0.0),
            "per_cell_residual": residuals}


def read_source(source: Path) -> dict:
    """Read and verify all cells, including unused sections and trailing zeros."""
    from openpyxl import load_workbook

    verify_file(source, SOURCE)
    formula_book = load_workbook(source, read_only=True, data_only=False, keep_links=False)
    value_book = load_workbook(source, read_only=True, data_only=True, keep_links=False)
    try:
        if formula_book.sheetnames != ["Sheet1"] or value_book.sheetnames != ["Sheet1"]:
            raise ValueError("Unexpected Figure 7 sheets")
        sheet, cached = formula_book.active, value_book.active
        if (sheet.max_row, sheet.max_column) != (2932, 16):
            raise ValueError("Unexpected Figure 7 dimensions")
        rows, values, formulas = [], {}, {}
        for raw_row, cached_row in zip(sheet, cached, strict=True):
            result = []
            for raw, cell in zip(raw_row, cached_row, strict=True):
                value = cell.value
                if raw.data_type == "e" or cell.data_type == "e":
                    raise ValueError(f"Source error at {raw.coordinate}")
                if value is not None and type(value) not in (int, float, str):
                    raise ValueError(f"Unexpected value at {cell.coordinate}")
                if isinstance(value, float) and not math.isfinite(value):
                    raise ValueError(f"Nonfinite source at {cell.coordinate}")
                if value is not None:
                    values[cell.coordinate] = value
                if raw.data_type == "f":
                    formulas[raw.coordinate] = raw.value
                elif raw.value != value:
                    raise ValueError(f"Nonformula cell changed at {raw.coordinate}")
                result.append(value)
            rows.append(result)
        audit = check_formulas(values, formulas)
        if audit["formula_count"] != 888:
            raise ValueError("Unexpected source formula count")
        return {"schema": 1, "source": SOURCE, "sheet": "Sheet1", "rows": rows,
                "formulas": formulas, "formula_audit": audit}
    finally:
        formula_book.close()
        value_book.close()


def sections(data: dict) -> dict:
    rows = data["rows"]
    expected = {"A1": "DL5, firing rate (spikes/s)", "A4": "I (pA)",
                "A917": "DL5 PN, uEPSC (pA)", "A920": "time (ms)",
                "A2924": "VA6, depolarization (mV)"}
    for address, title in expected.items():
        if rows[int(address[1:])-1][0] != title:
            raise ValueError(f"Changed section {address}")
    if rows[1][1] != "E2-hexenal" or rows[1][6] != "solvent":
        raise ValueError("Changed ramp conditions")
    if (rows[917][1], rows[917][9], rows[918][1], rows[918][9]) != (
            "solvent", "E2-hexenal", 7, 5):
        raise ValueError("Changed uEPSC conditions or sample sizes")
    ramp = np.array([[number(r[c]) for c in (0, 1, 2, 3, 4, 6, 7, 8, 9)]
                     for r in rows[4:914]])
    # This abrupt return to zero is retained, not interpreted as a second trial.
    increasing = ramp[:890]
    if not np.all(np.diff(increasing[:, 0]) > 0) or np.any(ramp[890:] != 0):
        raise ValueError("Published ramp/tail structure changed")
    epsc = np.array([[number(r[c]) for c in (0, 1, 2, 3, 4, 5, 6, 7, 9, 10, 11, 12, 13)]
                     for r in rows[920:2921]])
    return {"ramp": ramp, "epsc": epsc}


def waveform_metrics(time_ms, current_pa, *, baseline_end_ms=20.0) -> dict:
    """Signed inward current, no rectification or unit conversion.

    Baseline and 50/100 ms integration windows are our declared analysis, not
    the authors' exact estimator. Samples are not independent biological trials.
    """
    t, y = np.asarray(time_ms, dtype=float), np.asarray(current_pa, dtype=float)
    if t.ndim != 1 or y.shape != t.shape or len(t) < 3:
        raise ValueError("Expected equal one-dimensional waveform arrays")
    if not np.isfinite(t).all() or not np.isfinite(y).all() or not np.all(np.diff(t) > 0):
        raise ValueError("Invalid waveform values or time ordering")
    baseline_mask = t < baseline_end_ms
    if not baseline_mask.any():
        raise ValueError("No baseline samples")
    baseline = float(y[baseline_mask].mean())
    peak = int(y.argmin())
    if t[peak] <= baseline_end_ms or t[-1] < t[peak] + 100:
        raise ValueError("Insufficient baseline or postpeak support")
    inward = baseline - y[peak:]
    post_t = t[peak:] - t[peak]
    if inward[0] <= 0:
        raise ValueError("No inward response")

    def integrate(start, end):
        inside = (post_t > start) & (post_t < end)
        tt = np.r_[start, post_t[inside], end]
        yy = np.interp(tt, post_t, inward)
        return float(np.sum(np.diff(tt) * (yy[:-1] + yy[1:]) / 2))

    below = np.flatnonzero(inward <= inward[0] / np.e)
    if len(below) == 0:
        raise ValueError("No 1/e crossing")
    k = int(below[0])
    crossing = float(np.interp(inward[0]/np.e, inward[k-1:k+1][::-1], post_t[k-1:k+1][::-1]))
    windows = {}
    for end in (50, 100):
        area = integrate(0, end)
        if area <= 0:
            raise ValueError("Nonpositive net inward charge")
        windows[str(end)] = {"postpeak_charge_pa_ms": area,
                             "charge_after_5ms_fraction": integrate(5, end)/area}
    return {"baseline_pa": baseline, "peak_source_time_ms": float(t[peak]),
            "peak_inward_pa": float(inward[0]), "first_1_over_e_crossing_ms": crossing,
            "postpeak_windows_ms": windows,
            "baseline_end_ms": baseline_end_ms}


def analyze(data: dict) -> dict:
    arrays = sections(data)
    epsc = arrays["epsc"]
    measurements = []
    columns = [("solvent", c) for c in "BCDEFGH"] + [("E2-hexenal", c) for c in "JKLMN"]
    for i, (condition, column) in enumerate(columns, 1):
        measurements.append({"condition": condition, "source_column": column,
                             "source_rows": [921, 2921],
                             "measurement_unit": "one PN per fly, peak-aligned trial-average uEPSC",
                             **waveform_metrics(epsc[:, 0], epsc[:, i]),
                             "baseline_sensitivity": [
                                 waveform_metrics(epsc[:, 0], epsc[:, i], baseline_end_ms=end)
                                 for end in (30.0, 40.0)]})
    ramp = arrays["ramp"]
    return {"schema": 1, "source": SOURCE, "formula_audit": data["formula_audit"],
            "epsc": measurements,
            "ramp_audit": {"source_rows": [5, 914], "rows": len(ramp),
                           "increasing_rows": 890, "trailing_all_zero_rows": 20,
                           "max_source_current_pa": float(ramp[:, 0].max()),
                           "source_step_pa": float(np.median(np.diff(ramp[:890, 0]))),
                           "n_cells_per_condition": 4,
                           "raw_spike_times_available": False,
                           "step_protocol_table_present": False,
                           "status": "No separate 5-pA current-step series in this workbook. Do not relabel ramp bins as step responses or independent cells."},
            "limits": [
                "Current is from voltage-clamp recordings, not calcium, membrane voltage or spike amplitude.",
                "Peak-aligned averages cannot determine cleft delay, optogenetic latency or trial jitter.",
                "DL5 intrinsic current response does not calibrate all PNs, KCs or APL.",
                "No equivalence test or chronic-exposure plasticity reproduction is claimed.",
                "VA6 lateral odor measurements are retained, not assigned to isolated direct ORN drive.",
            ]}


def extract(source: Path, output: Path) -> dict:
    if output.exists():
        raise FileExistsError(output)
    data = read_source(source)
    result = analyze(data)
    output.mkdir(parents=True)
    dump_new(output / "cells.json", data)
    dump_new(output / "analysis.json", result)
    dump_new(output / "extraction.json", {"source": SOURCE,
                                          "cells_sha256": digest(output / "cells.json")})
    return result


def native_input_probe(graph_path: Path, output: Path, *, ticks: int = 160) -> dict:
    """A release imposed at one anatomical ORN receptor, not a simulated ORN.

    Isolate an already selected PN, preserving every incident source row and
    address. Read native somatic integration to separate arriving effective
    current from lingering voltage. Physical clock and pA scale remain unset.
    """
    from .connectome import Subgraph
    from .paula import Dynamics, build_paula

    if output.exists():
        raise FileExistsError(output)
    if type(ticks) is not int or ticks < 16:
        raise ValueError("Need at least 16 ticks")
    graph = Subgraph.load(graph_path)
    candidates = [r for r in graph.selected if graph.nodes[r]["annotation"]["hemibrain_type"] == "DL5_adPN"]
    if len(candidates) != 1:
        raise ValueError("Expected one selected DL5_adPN")
    root = candidates[0]
    edges = graph.edges[(graph.edges[:, 0] == int(root)) | (graph.edges[:, 1] == int(root))]
    endpoints = {str(v) for v in edges[:, :2].flat} | {root}
    isolated = Subgraph((root,), {r: graph.nodes[r] for r in endpoints}, edges,
                        {"parent_provenance": graph.provenance, "cut": "one DL5 PN"})
    dynamics = Dynamics(weight_per_count=0.075)
    prep = build_paula(isolated, dynamics)
    cell = prep.network.network.neurons[prep.root_to_id[root]]
    orn_edges = sorted((e for e in edges if str(e[1]) == root and
                        graph.nodes[str(e[0])]["annotation"]["hemibrain_type"] == "ORN_DL5"),
                       key=lambda e: int(e[8]))
    if not orn_edges:
        raise ValueError("No identified ORN_DL5 input")
    edge = orn_edges[0]
    binding = prep.incoming_boundary_ports[prep.incoming_boundary_ports[:, 0] == edge[8]]
    if binding.shape != (1, 4):
        raise ValueError("Ambiguous incoming receptor binding")
    sid = int(binding[0, 3])
    delay = cell.distances[sid]
    initial_weight = float(cell.postsynaptic_points[sid].u_i.info)
    trace = []
    for tick in range(ticks):
        if tick == 0:
            cell.input_buffer[sid, 0] = 1.0
        before = float(cell.S)
        queue = list(cell.propagation_queue)
        cell.tick({}, tick)
        if cell.O != 0:
            raise ValueError("Unexpected spike invalidates the subthreshold current inversion")
        inferred = dynamics.lambda_ticks * (float(cell.S)-before) + before
        # Independent queue readout. The receptor event at tick zero is not in
        # the pre-tick queue; its zero-delay case is handled explicitly.
        arrived = sum((v * dynamics.signal_decay**cell.distances[p]
                       for at, _, v, p in queue if at <= tick), 0.0)
        if tick == 0 and delay == 0:
            arrived += initial_weight
        trace.append([tick, before, float(cell.S), inferred, arrived,
                      float(cell.postsynaptic_points[sid].u_i.info), float(cell.t_ref)])
    trace = np.asarray(trace)
    # Somatic arithmetic can be float32. Inverting its subtraction amplifies
    # rounding by lambda; do not claim bit-exact inferred current.
    current_atol = 8 * np.finfo(np.float32).eps * max(1.0, np.abs(trace[:, 4]).max())
    if not np.isfinite(trace).all() or not np.allclose(trace[:, 3], trace[:, 4], atol=current_atol, rtol=0):
        raise ValueError("Native current and independent queue readout disagree")
    result = {"schema": 1, "root": root, "global_index": cell.id,
              "source_hashes": {str(p): digest(p) for p in (
                  graph_path / "manifest.json", Path(__file__),
                  Path(__file__).with_name("paula.py"),
                  Path(inspect.getfile(type(cell))))},
              "source_row": int(edge[8]), "orn_root": str(edge[0]),
              "counted_contacts": int(edge[4]), "postsynaptic_port": sid,
              "num_inputs_with_experimental_port": cell.params.num_inputs,
              "preserved_incoming_pairs": len(prep.incoming_boundary_ports),
              "preserved_outgoing_pairs": len(prep.outgoing_boundary_terminals),
              "parameters": prep.assumptions,
              "stimulus": "unit release at receptor tick 0; upstream ORN not simulated",
              "effective_current_nonzero_ticks": trace[trace[:, 4] != 0, 0].astype(int).tolist(),
              "arrived_current_peak": float(trace[:, 4].max()),
              "maximum_current_reconstruction_error": float(np.max(np.abs(trace[:, 3]-trace[:, 4]))),
              "current_reconstruction_tolerance": float(current_atol),
              "weight_initial": initial_weight, "weight_final": float(trace[-1, 5]),
              "physical_seconds_per_tick": None, "paula_current_to_pa": None,
              "claim": "Native receptor-to-hillock kernel witness only, not a voltage-clamp simulation or biological replication"}
    output.mkdir(parents=True)
    with (output / "per_tick.npz").open("xb") as stream:
        np.savez_compressed(stream, trace=trace, columns=np.array([
            "tick", "S_before", "S_after", "current_from_S", "current_from_queue", "post_weight", "t_ref"]))
    dump_new(output / "analysis.json", result)
    return result


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    ex = sub.add_parser("extract")
    ex.add_argument("workbook", type=Path)
    ex.add_argument("output", type=Path)
    probe = sub.add_parser("native-input")
    probe.add_argument("graph", type=Path)
    probe.add_argument("output", type=Path)
    args = parser.parse_args()
    result = (extract(args.workbook, args.output) if args.command == "extract" else
              native_input_probe(args.graph, args.output))
    print(json.dumps({k: v for k, v in result.items() if k not in ("formula_audit", "parameters")}, indent=2))


if __name__ == "__main__":
    main()
