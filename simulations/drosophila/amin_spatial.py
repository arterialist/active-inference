"""Spatial observation benchmark for Amin et al. 2020, not an odor experiment.

Reproduce the authors' straight-backbone exponential calculation, read their
published connectome predictions separately, and test a conservative cable on
the same schematic backbone. The backbone is NOT the APL neurite skeleton.
No fitted calcium-to-voltage conversion, neuronal learning or fly acceptance
is inferred. All 35 observed segments count once, including the two distinct
segments that the schematic places at identical coordinates.
"""
from __future__ import annotations

import argparse
import hashlib
import json
from pathlib import Path

SOURCES = {
    "elife-56954-fig7-data1-v2.xlsx": "ba1d7c4dfaa2007aa3e71f668b72ee81845ea20094012f44450117a3a2165e9c",
    "elife-56954-fig8-data1-v2.xlsx": "673b933de6b2aaa4382d7e857bfb45154fff21766fbc826fea750465a10986b3",
    "redDyeVoronoiFits.mat": "249626c5f3abe0f57053124025d75b0a7e8c338d4e7f1aa381d4f6d432dcf155",
    "skel2d.mat": "72c765cc9d790ffb90eed0e4f78a721f2c67b822ecb0812abcb2413b9f797453",
    "voronoiIndices.mat": "a5974840298f66d53554bf4516a33c5ce6cc36120ff590f79978db3b6d654686",
}
AUTHOR_COMMIT = "d16f97f26e605ec0591043db636b7ff8e9801e0f"
ARTICLE = "https://doi.org/10.7554/eLife.56954"
CELLS_SHA256 = "d5b3072dfc2c5bc25eebffc0d8c2488c160b4ac08987f66eb5da807796b0191f"


def digest(path):
    with Path(path).open("rb") as stream:
        return hashlib.file_digest(stream, "sha256").hexdigest()


def write_new(path, value):
    with Path(path).open("x") as stream:
        json.dump(value, stream, indent=2, allow_nan=False)
        stream.write("\n")


def verify(directory):
    for name, expected in SOURCES.items():
        if digest(directory / name) != expected:
            raise ValueError(f"Primary source differs: {name}")


def extract(directory, output):
    """Read originals, including formulas and their saved values, without export.

    Run using the bundled document Python. Formula caches are preserved as
    caches, not claimed to have been recalculated. Analysis below uses numeric
    source observations and checks the dye against independent MAT arrays.
    """
    from openpyxl import load_workbook

    if output.exists():
        raise FileExistsError(output)
    verify(directory)
    result = {"sources": SOURCES, "workbooks": {}}
    for name in SOURCES:
        if not name.endswith(".xlsx"):
            continue
        values = load_workbook(directory / name, read_only=True, data_only=True, keep_links=False)
        formulas = load_workbook(directory / name, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = {}
            for s, f in zip(values, formulas, strict=True):
                if s.title != f.title:
                    raise ValueError("Workbook views differ")
                rows, expressions = [], {}
                for vr, fr in zip(s, f, strict=True):
                    row = []
                    for v, original in zip(vr, fr, strict=True):
                        if original.data_type == "e":
                            raise ValueError(f"Source error: {name}:{s.title}:{original.coordinate}")
                        if original.data_type == "f":
                            expressions[original.coordinate] = original.value
                        row.append(v.value)
                    rows.append(row)
                sheets[s.title] = {"rows": rows, "formulas": expressions}
            result["workbooks"][name] = sheets
        finally:
            values.close()
            formulas.close()
    write_new(output, result)
    return {"cells_sha256": digest(output), "source_hashes": SOURCES}


def finite(values):
    import numpy as np
    if any(type(v) not in {int, float} for v in values):
        raise ValueError("Missing or non-numeric source observation")
    result = np.array(values, dtype=float)
    if not np.isfinite(result).all():
        raise ValueError("Nonfinite source observation")
    return result


def merge_profiles(horizontal, vertical, hi, vi):
    """Remove duplicated common-stem observations, not distinct source regions."""
    import numpy as np
    result = np.full(35, np.nan)
    for values, indices in ((horizontal, hi), (vertical, vi)):
        if len(values) != len(indices):
            raise ValueError("Profile/index length mismatch")
        for v, index in zip(values, indices, strict=True):
            if not 0 <= index < 35:
                raise ValueError("Invalid Voronoi index")
            if np.isfinite(result[index]) and result[index] != v:
                raise ValueError("Shared source segment differs between branches")
            result[index] = v
    if not np.isfinite(result).all():
        raise ValueError("Not all 35 source segments are represented")
    return result


def exponential_profile(xy, drive, length):
    """The published city-block-distance equation, independently evaluated.

    Reference algorithm: authors' convolveExpDecay.m, pinned AUTHOR_COMMIT.
    Equal stimulus-site contributions are part of that reference, not a
    membrane-area/conductance assumption for the PAULA cable.
    """
    import numpy as np
    xy, drive = np.asarray(xy), np.asarray(drive)
    if (xy.ndim != 2 or xy.shape[1] != 2 or drive.shape != (len(xy),)
            or not np.isfinite(xy).all() or not np.isfinite(drive).all()
            or np.any(drive < 0) or not np.any(drive > 0)):
        raise ValueError("Finite geometry and nonnegative nonzero stimulus required")
    if not np.isfinite(length) or length <= 0:
        raise ValueError("Positive finite space constant required")
    distance = np.abs(xy[:, None, :] - xy[None, :, :]).sum(axis=2)
    result = np.exp(-distance / length) @ drive
    return result / result.max()


def backbone_cable(xy, drive, length, ticks=400):
    """Explicit geometry ablation, not a reconstructed neuron or calcium model.

    Merge the two coincident schematic sites into one physical junction and
    average their drive there. Uniform 0.2 um radius, sealed ends and uniform
    membrane drive density are assumptions. All source sites remain readable.
    A continuous held stimulus tests the stationary operator, not ATP timing.
    """
    import numpy as np
    from scipy.sparse import diags
    from scipy.sparse.linalg import spsolve
    from simulations.paula_loader import ensure_paula_available
    ensure_paula_available()
    from neuron.extensions.experimental.passive_cable import PassiveCable

    xy, drive = np.asarray(xy), np.asarray(drive)
    if (not np.isfinite(length) or length <= 0 or type(ticks) is not int or ticks < 1
            or xy.ndim != 2 or xy.shape[1] != 2 or drive.shape != (len(xy),)
            or not np.isfinite(xy).all() or not np.isfinite(drive).all()
            or np.any(drive < 0) or not np.any(drive > 0)):
        raise ValueError("Invalid schematic cable stimulus, geometry or duration")
    positions, inverse = np.unique(xy, axis=0, return_inverse=True)
    parents = np.full(len(positions), -1, dtype=np.int64)
    for i, p in enumerate(positions):
        if np.all(p == 0):
            continue
        q = p - 10 * np.sign(p)
        candidates = np.flatnonzero(np.all(positions == q, axis=1))
        if len(candidates) != 1:
            raise ValueError("Backbone is not a connected 10 um axis-aligned tree")
        parents[i] = candidates[0]
    radius = .2
    cable = PassiveCable(parents, np.column_stack((positions, np.zeros(len(positions)))),
                         np.full(len(positions), radius), 2 * length**2 / radius)
    density = np.bincount(inverse, weights=drive) / np.bincount(inverse)
    current = cable.capacity * density
    stationary = spsolve(diags(cable.capacity) + cable.laplacian, current)
    trajectory = [cable.voltage.copy()]
    residuals = []
    for _ in range(ticks):
        old = cable.voltage.copy()
        cable.step(current, .05)
        residuals.append(float(np.max(abs(cable.capacity * cable.voltage
            + .05 * (cable.laplacian @ cable.voltage) - .95 * cable.capacity * old - .05 * current))))
        trajectory.append(cable.voltage.copy())
    error = float(np.max(abs(cable.voltage - stationary)))
    if error > 1e-8 or max(residuals) > 1e-12:
        raise ValueError("Schematic cable did not reach its checked stationary solution")
    return {"positions": positions, "parents": parents, "source_to_node": inverse,
            "tick": np.arange(ticks + 1), "alpha": .05, "radius_um": radius,
            "space_constant_um": length, "rm_over_ra_um": 2 * length**2 / radius,
            "capacity": cable.capacity, "drive_density": density,
            "current": current, "voltage": np.array(trajectory),
            "stationary": stationary, "equation_residual": np.array(residuals),
            "stationary_error": error, "profile": (stationary / stationary.max())[inverse]}


def analyze(directory, cells_path, output):
    import inspect
    import numpy as np
    from scipy.io import loadmat

    if output.exists():
        raise FileExistsError(output)
    verify(directory)
    if digest(cells_path) != CELLS_SHA256:
        raise ValueError("Extracted cells differ from the verified original extraction")
    cells = json.loads(cells_path.read_text())
    if cells["sources"] != SOURCES:
        raise ValueError("Extracted source hashes differ")
    book = cells["workbooks"]["elife-56954-fig8-data1-v2.xlsx"]
    idx = loadmat(directory / "voronoiIndices.mat")
    hi, vi = (idx[name].ravel().astype(int) - 1 for name in ("ctohindices", "ctovindices"))
    xy = loadmat(directory / "skel2d.mat")["skel2d"].astype(float)
    dyes = loadmat(directory / "redDyeVoronoiFits.mat")
    if xy.shape != (35, 2) or len(hi) != 26 or len(vi) != 28:
        raise ValueError("Unexpected source geometry")
    rows = book["voronoi indices"]["rows"]
    if not np.array_equal(finite([r[10] for r in rows[3:29]]), hi + 1):
        raise ValueError("Workbook/MAT horizontal indices differ")
    if not np.array_equal(finite([r[11] for r in rows[3:31]]), vi + 1):
        raise ValueError("Workbook/MAT vertical indices differ")
    for column, key in ((4, "vstimFit"), (5, "hstimFit"), (6, "cstimFit")):
        if not np.array_equal(finite([r[column] for r in rows[3:38]]), dyes[key].ravel()):
            raise ValueError("Workbook/MAT dye fits differ")
    output.mkdir(parents=True)
    results, trace_files = [], []
    for site in ("H", "V", "C"):
        table = book[site + "stim"]["rows"]
        def profile(hcolumn, vcolumn):
            return merge_profiles(finite([r[hcolumn] for r in table[4:30]]),
                                  finite([r[vcolumn] for r in table[4:32]]), hi, vi)
        observed, measured_dye = profile(1, 3), profile(2, 4)
        fitted_dye = dyes[site.lower() + "stimFit"].ravel()
        candidates = {"uniform_global": np.ones(35), "fitted_dye_only": fitted_dye,
                      "measured_dye_only": measured_dye}
        reproductions = []
        for index, length in enumerate((25., 50., 75.)):
            conn_start, back_start = (13, 20) if site == "H" else (8, 15)
            published_backbone = profile(back_start + 2*index, back_start + 2*index + 1)
            reproduced = exponential_profile(xy, fitted_dye, length)
            difference = float(np.max(abs(reproduced - published_backbone)))
            # Numerical reproduction tolerance, not a physiological fit bar.
            if difference > 1e-12:
                raise ValueError(f"Published backbone curve not reproduced: {site}, {length}, {difference}")
            reproductions.append({"length_um": length, "max_abs_error": difference})
            candidates[f"published_connectome_{length:g}"] = profile(conn_start + 2*index, conn_start + 2*index + 1)
            candidates[f"reproduced_backbone_kernel_{length:g}"] = reproduced
            cable = backbone_cable(xy, fitted_dye, length)
            candidates[f"schematic_conservative_cable_{length:g}"] = cable.pop("profile")
            trace = output / f"{site.lower()}-cable-{length:g}.npz"
            with trace.open("xb") as stream:
                np.savez_compressed(stream, **cable)
            trace_files.append({"file": trace.name, "sha256": digest(trace)})
        errors = {key: float(np.sqrt(np.mean((value - observed)**2))) for key, value in candidates.items()}
        results.append({"site": site, "observed": observed.tolist(),
            "measured_dye": measured_dye.tolist(), "fitted_dye": fitted_dye.tolist(),
            "predicted_profiles": {k: v.tolist() for k, v in candidates.items()},
            "rmse_35_unique_source_regions": errors, "numerical_reproduction": reproductions})
    aggregate = {key: float(np.sqrt(np.mean([r["rmse_35_unique_source_regions"][key]**2 for r in results])))
                 for key in results[0]["rmse_35_unique_source_regions"]}
    uniform_checks = []
    for length in (25., 50., 75.):
        kernel = exponential_profile(xy, np.ones(35), length)
        cable = backbone_cable(xy, np.ones(35), length)
        # Inspect the entire trajectory, not only its converged endpoint.
        max_range = float(np.ptp(cable["voltage"], axis=1).max())
        if max_range > 1e-12:
            raise ValueError("Uniform membrane-density input became spatially nonuniform")
        uniform_checks.append({"length_um": length, "kernel_normalized_min": float(kernel.min()),
            "kernel_normalized_max": float(kernel.max()), "cable_max_spatial_range_any_tick": max_range})
        trace = output / f"uniform-cable-{length:g}.npz"
        with trace.open("xb") as stream:
            np.savez_compressed(stream, **cable)
        trace_files.append({"file": trace.name, "sha256": digest(trace)})
    result = {"schema": "amin-spatial-observation-v1", "article": ARTICLE,
        "author_code_commit": AUTHOR_COMMIT, "original_sha256": SOURCES,
        "cells_sha256": digest(cells_path), "analysis_source_sha256": digest(__file__),
        "regions_per_site": 35, "sites": results, "rmse_equal_sites_and_regions": aggregate,
        "uniform_density_diagnostic": uniform_checks, "traces": trace_files,
        "limits": ["Published normalized spatial means, not raw imaging trajectories or independent flies",
                   "RMSE is a descriptive equal-region comparison, not uncertainty, significance or acceptance",
                   "Connectome curves read from the paper, not independently rerun",
                   "Schematic cable uses the imaging backbone, not hemibrain or FlyWire APL morphology",
                   "400 model ticks of held drive, not the measured ATP time course",
                   "No calcium/voltage mapping, no sensory coding or learning claim",
                   "No changed PAULA neuron equations or frozen network weights"]}
    from neuron.extensions.experimental.passive_cable import PassiveCable
    result["cable_source_sha256"] = digest(inspect.getfile(PassiveCable))
    write_new(output / "analysis.json", result)
    return {"numerical_reproductions": 9, "rmse": aggregate, "trace_files": len(trace_files)}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    ex = sub.add_parser("extract")
    ex.add_argument("directory", type=Path)
    ex.add_argument("output", type=Path)
    an = sub.add_parser("analyze")
    an.add_argument("directory", type=Path)
    an.add_argument("cells", type=Path)
    an.add_argument("output", type=Path)
    args = parser.parse_args()
    result = extract(args.directory, args.output) if args.command == "extract" else analyze(args.directory, args.cells, args.output)
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
