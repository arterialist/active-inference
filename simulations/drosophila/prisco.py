"""Read-only import and analysis of Prisco et al. primary physiological tables.

No input currents are inferred. Glomerulus labels, active optical ROIs and
reconstructed neurons are different experimental units. Extraction uses
openpyxl read-only; the analysis command needs only the simulation dependencies.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import math
from pathlib import Path
from statistics import mean

SOURCES = Path(__file__).with_name("prisco2021_sources.json")
GLOMERULI = ("DM6", "DC2", "VA6", "VA1", "DA1", "DL1", "DL5", "DM3", "DC1")


def digest(path: Path) -> str:
    with path.open("rb") as f:
        return hashlib.file_digest(f, "sha256").hexdigest()


def dump_new(path: Path, value) -> None:
    with path.open("x") as f:
        json.dump(value, f, ensure_ascii=False, allow_nan=False, indent=2)
        f.write("\n")


def verify_file(path: Path, spec: dict) -> None:
    if path.stat().st_size != spec["bytes"] or digest(path) != spec["sha256"]:
        raise ValueError(f"Primary source bytes differ: {path}")


def extract(source: Path, output: Path) -> dict:
    """Verify the public repository hashes, preserve originals and every cell.

    Use the bundled document runtime for this command. No download, recalculation,
    workbook edit, macro execution or inferred animal identity occurs here.
    """
    from openpyxl import load_workbook

    if output.exists():
        raise FileExistsError(output)
    manifest = json.loads(SOURCES.read_text())
    books = {}
    for spec in manifest["files"]:
        filename = spec["filename"]
        if Path(filename).name != filename:
            raise ValueError("Source filename is not a basename")
        path = source / filename
        verify_file(path, spec)
        if not filename.endswith(".xlsx"):
            continue
        wb = load_workbook(path, read_only=True, data_only=False, keep_links=False)
        try:
            sheets = []
            for sheet in wb:
                rows = []
                for row in sheet:
                    values = []
                    for cell in row:
                        if cell.data_type in {"f", "e"}:
                            raise ValueError(f"Unexpected formula/error: {filename}:{cell.coordinate}")
                        value = cell.value
                        if value is not None and type(value) not in {str, int, float}:
                            raise ValueError(f"Unexpected cell type: {filename}:{cell.coordinate}")
                        if isinstance(value, float) and not math.isfinite(value):
                            raise ValueError(f"Nonfinite cell: {filename}:{cell.coordinate}")
                        values.append(value)
                    rows.append(values)
                sheets.append({"name": sheet.title, "rows": rows,
                               "dimensions": [sheet.max_row, sheet.max_column]})
            books[str(spec["id"])] = {"source": spec, "sheets": sheets}
        finally:
            wb.close()
    # All inputs validate before creating any output. Copies are byte-for-byte,
    # not re-exported workbooks, and never replace a user's existing file.
    output.mkdir(parents=True)
    for spec in manifest["files"]:
        with (output / spec["filename"]).open("xb") as target:
            target.write((source / spec["filename"]).read_bytes())
        verify_file(output / spec["filename"], spec)
    result = {"schema": 1, "sources": manifest, "workbooks": books,
              "units": "As published; no rescaling or missing-value imputation",
              "animal_identity": "Local row positions only, not cross-workbook animal IDs"}
    dump_new(output / "cells.json", result)
    dump_new(output / "extraction.json", {"cells_sha256": digest(output / "cells.json"),
                                           "source_manifest_sha256": digest(SOURCES)})
    return {"workbooks": len(books), "source_bytes": sum(f["bytes"] for f in manifest["files"])}


def load_cells(directory: Path) -> dict:
    check = json.loads((directory / "extraction.json").read_text())
    if digest(directory / "cells.json") != check["cells_sha256"]:
        raise ValueError("Extracted cells changed")
    if digest(SOURCES) != check["source_manifest_sha256"]:
        raise ValueError("Pinned source manifest changed")
    data = json.loads((directory / "cells.json").read_text())
    for spec in json.loads(SOURCES.read_text())["files"]:
        verify_file(directory / spec["filename"], spec)
    return data


def sheet(data: dict, file_id: int) -> tuple[str, list]:
    sheets = data["workbooks"][str(file_id)]["sheets"]
    if len(sheets) != 1:
        raise ValueError(f"Unexpected sheet count: {file_id}")
    return sheets[0]["name"], sheets[0]["rows"]


def number(value) -> float:
    if type(value) not in {int, float} or not math.isfinite(value):
        raise ValueError(f"Missing/non-numeric observation: {value!r}")
    return float(value)


def pn_observations(data: dict, file_id: int) -> list[dict]:
    name, rows = sheet(data, file_id)
    first_odor = {1282866: "δ-DL", 1282868: "Mch"}[file_id]
    if rows[0] != [v for g in GLOMERULI for v in (g, None)]:
        raise ValueError("Glomerulus headers changed")
    if rows[1] != [v for _ in GLOMERULI for v in (first_odor, "Oct")]:
        raise ValueError("Odor headers changed")
    if len(rows) != 12 or any(len(r) != 18 for r in rows):
        raise ValueError("Expected ten rows and nine glomerulus pairs")
    result = []
    for row_index, row in enumerate(rows[2:], 3):
        for i, glomerulus in enumerate(GLOMERULI):
            for offset, odor in enumerate((first_odor, "Oct")):
                col = 2 * i + offset
                result.append({"cohort": str(file_id), "source_row": row_index,
                               "sheet": name, "cell": f"{chr(65 + col)}{row_index}",
                               "glomerulus": glomerulus, "odor": odor,
                               "peak_dff_percent": number(row[col])})
    return result


def glomerulus_candidates(nodes: dict, selected: tuple[str, ...]) -> dict:
    """Exact uniglomerular label matches only, not a physiological assignment.

    VA1 is explicitly ambiguous; VA1d/VA1v are not repaired via prefix matching.
    Even an exact label does not prove GH146 expression, identical excitability,
    or a measurement for an individual reconstructed cell.
    """
    candidates = {g: [] for g in GLOMERULI}
    unresolved = []
    for root in selected:
        node = nodes[root]
        a = node["annotation"]
        if a["cell_class"] != "ALPN":
            continue
        record = {"root": root, "global_index": node["global_index"],
                  "type": a.get("hemibrain_type", ""), "side": a.get("side", ""),
                  "subclass": a.get("cell_sub_class", ""),
                  "known_nt": a.get("known_nt", ""),
                  "top_nt": a.get("top_nt", ""),
                  "driver_expression_verified": False}
        label, sep, _ = record["type"].partition("_")
        if (sep and record["subclass"] == "uniglomerular" and
                label in candidates and label != "VA1" and "," not in record["type"]):
            candidates[label].append(record)
        else:
            record["reason"] = ("VA1 subtype unresolved" if label in {"VA1", "VA1d", "VA1v"}
                                else "No exact measured uniglomerular label")
            unresolved.append(record)
    return {"candidate_neurons": candidates, "unresolved": unresolved,
            "status": "Label coverage only; no per-neuron calcium or current assigned",
            "candidate_count": sum(map(len, candidates.values())),
            "selected_alpn_count": sum(map(len, candidates.values())) + len(unresolved)}


def paired_peaks(data: dict, file_id: int, start: int, left: int, right: int,
                 odors: tuple[str, str], condition: str, expected_n: int) -> dict:
    from scipy.stats import ttest_rel

    name, rows = sheet(data, file_id)
    values = []
    for r, row in enumerate(rows[start - 1:], start):
        if all(v is None for v in row):
            continue
        x, y = number(row[left]), number(row[right])
        values.append({"source_row": r, "left_cell": f"{chr(65+left)}{r}",
                       "right_cell": f"{chr(65+right)}{r}",
                       "first_peak": x, "second_peak": y, "second_minus_first": y-x})
    if len(values) != expected_n:
        raise ValueError(f"Unexpected animal rows: {file_id}")
    x = [v["first_peak"] for v in values]
    y = [v["second_peak"] for v in values]
    test = ttest_rel(y, x)
    return {"source_file_id": file_id, "sheet": name, "condition": condition,
            "odors": odors, "units": "peak ΔF/F0 percent",
            "n_animal_rows": len(values), "animals": values,
            "paired_t_reanalysis": {"statistic": float(test.statistic), "p": float(test.pvalue)},
            "mean_first": mean(x), "mean_second": mean(y),
            "mean_second_minus_first": mean(v["second_minus_first"] for v in values),
            "status": "Reanalysis of row-paired peaks, not equivalence or model acceptance"}


def pooled_reconciliation(data: dict, summary_id: int, pooled_id: int,
                          start: int, mean_col: int, count_col: int,
                          pooled_start: int, pooled_col: int) -> dict:
    """No individual ROI-to-animal assignment is recoverable from these columns."""
    _, rows = sheet(data, summary_id)
    _, pooled_rows = sheet(data, pooled_id)
    summaries = [(number(r[mean_col]), number(r[count_col])) for r in rows[start-1:]
                 if any(v is not None for v in r)]
    if any(n < 1 or n != int(n) for _, n in summaries):
        raise ValueError("Invalid active ROI count")
    pooled = [number(r[pooled_col]) for r in pooled_rows[pooled_start-1:]
              if r[pooled_col] is not None]
    if not pooled:
        raise ValueError("Empty pooled observations")
    total = int(sum(n for _, n in summaries))
    weighted = sum(v*n for v, n in summaries)/total
    return {"summary_file_id": summary_id, "pooled_file_id": pooled_id,
            "summary_mean_column": mean_col + 1, "pooled_column": pooled_col + 1,
            "n_animals": len(summaries), "sum_reported_active_rois": total,
            "pooled_roi_values": len(pooled), "counts_match": total == len(pooled),
            "mean_from_reported_counts": weighted, "mean_of_pooled_peaks": mean(pooled),
            "difference": mean(pooled)-weighted,
            "roi_to_animal_identity_available": False,
            "status": "Consistency check only; never infer missing ROI values or group membership"}


def analyze(data: dict, graph=None) -> dict:
    from scipy.stats import t, ttest_ind
    from statistics import variance

    pn = pn_observations(data, 1282866) + pn_observations(data, 1282868)
    comparisons = [paired_peaks(data, *args) for args in (
        (1282858, 2, 0, 1, ("Mch", "Oct"), "APL calyx", 10),
        (1282857, 2, 0, 1, ("δ-DL", "Oct"), "APL calyx", 10),
        (1282867, 2, 0, 1, ("Mch", "Oct"), "PN AL average", 10),
        (1282865, 2, 0, 1, ("δ-DL", "Oct"), "PN AL average", 10),
        (1282876, 3, 0, 3, ("Mch", "Oct"), "Active PN boutons", 10),
        (1282875, 3, 0, 3, ("Mch", "Oct"), "Active KC claws", 9),
        (1282870, 3, 0, 3, ("δ-DL", "Oct"), "Active KC claws", 10),
        (1282863, 4, 0, 3, ("Mch", "Oct"), "Active KC claws APL ON", 10),
        (1282863, 4, 6, 9, ("Mch", "Oct"), "Active KC claws APL OFF", 10),
    )]
    consistency = [pooled_reconciliation(data, *args) for args in (
        (1282876, 1282874, 3, 0, 2, 2, 0), (1282876, 1282874, 3, 3, 5, 2, 1),
        (1282875, 1282873, 3, 0, 2, 2, 0), (1282875, 1282873, 3, 3, 5, 2, 1),
        (1282870, 1282871, 3, 0, 2, 2, 0), (1282870, 1282871, 3, 3, 5, 2, 1),
        (1282863, 1282864, 4, 0, 2, 3, 0), (1282863, 1282864, 4, 3, 5, 3, 1),
        (1282863, 1282864, 4, 6, 8, 3, 2), (1282863, 1282864, 4, 9, 11, 3, 3),
    )]
    # The source delta sheet is retained, including its sign labels. Check its
    # numbers rather than treating those labels as executable subtraction rules.
    delta_name, delta_rows = sheet(data, 1282869)
    delta_checks = []
    for col, comparison_index in enumerate((4, 5, 7, 8)):
        c = comparisons[comparison_index]
        residuals = [number(delta_rows[i+1][col])-a["second_minus_first"]
                     for i, a in enumerate(c["animals"])]
        delta_checks.append({"source_header": delta_rows[0][col], "sheet": delta_name,
                             "column": col+1, "numeric_comparison": "Oct minus Mch",
                             "per_row_residuals": residuals})
    # Two odors are paired within each animal, but APL ON and OFF animals are
    # separate groups. Never pair those groups just because both have ten rows.
    on = [a["second_minus_first"] for a in comparisons[7]["animals"]]
    off = [a["second_minus_first"] for a in comparisons[8]["animals"]]
    v_on, v_off = variance(on)/len(on), variance(off)/len(off)
    se = math.sqrt(v_on + v_off)
    df = (v_on+v_off)**2/(v_on**2/(len(on)-1)+v_off**2/(len(off)-1))
    difference = mean(off)-mean(on)
    half_width = float(t.ppf(.975, df))*se
    contrast = {"comparison": "APL OFF minus ON, within-animal Oct-minus-Mch peak difference",
                "units": "percentage points of peak ΔF/F0",
                "mean_difference": difference, "standard_error": se, "welch_df": df,
                "welch_95_percent_ci": [difference-half_width, difference+half_width],
                "welch_p": float(ttest_ind(off, on, equal_var=False).pvalue),
                "on_differences": on, "off_differences": off,
                "status": "Independent reanalysis, not the published ANOVA/Tukey test or an acceptance threshold"}
    paper_checks = []
    for index, figure, published_p in ((4, "3B", .0002), (5, "3F", .1648)):
        calculated = comparisons[index]["paired_t_reanalysis"]["p"]
        paper_checks.append({"figure": figure, "test": "two-sided paired t-test",
                             "source_file_id": comparisons[index]["source_file_id"],
                             "published_p": published_p, "recomputed_p": calculated,
                             "matches_to_published_decimals": round(calculated, 4) == published_p,
                             "status": "Reproduction check, not a biological acceptance criterion"})
    return {"schema": 1, "article_doi": "10.7554/eLife.74172",
            "source_manifest": data["sources"],
            "pn_observations": pn, "paired_peak_comparisons": comparisons,
            "pooled_consistency": consistency, "delta_sign_checks": delta_checks,
            "apl_odor_contrast": contrast,
            "paper_test_reproduction": paper_checks,
            "coverage": glomerulus_candidates(graph.nodes, graph.selected) if graph else None,
            "limitations": [
                "Peak tables contain no time series; they cannot validate latency or dynamics.",
                "Nine glomerulus labels are not a complete PN odor pattern. VA1 remains ambiguous.",
                "Candidate neuron labels do not verify driver expression or individual responses.",
                "Different indicators, optical regions and cohorts cannot be joined by row index.",
                "Active ROIs are selected separately for each odor; zero spikes is not normalization.",
                "Pooled ROI columns lack animal/ROI identities; they are not paired observations.",
                "N in Mean/SD/N blocks counts active optical ROIs, not independent animals.",
                "SD columns are preserved but not used as weights; their scaling is not established here.",
                "The paired t reanalyses are not the paper's ANOVA/Tukey or Wilcoxon analyses.",
                "A non-significant difference does not prove equality or equivalence.",
                "No calcium-to-current, voltage-to-calcium or seconds-to-ticks conversion is inferred.",
            ]}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest="command", required=True)
    ex = sub.add_parser("extract")
    ex.add_argument("source", type=Path)
    ex.add_argument("output", type=Path)
    an = sub.add_parser("analyze")
    an.add_argument("source", type=Path)
    an.add_argument("output", type=Path)
    an.add_argument("--graph", type=Path)
    args = parser.parse_args()
    if args.command == "extract":
        print(json.dumps(extract(args.source, args.output)))
    else:
        graph = None
        if args.graph:
            from .connectome import Subgraph
            graph = Subgraph.load(args.graph)
        report = analyze(load_cells(args.source), graph)
        report["extracted_cells_sha256"] = digest(args.source / "cells.json")
        report["graph_files"] = ({f: digest(args.graph / f) for f in ("manifest.json", "nodes.json", "edges.npz")}
                                 if graph else None)
        dump_new(args.output, report)
        print(json.dumps({"pn_observations": len(report["pn_observations"]),
                          "paired_comparisons": len(report["paired_peak_comparisons"]),
                          "count_disagreements": sum(not r["counts_match"] for r in report["pooled_consistency"]),
                          "candidate_pns": report["coverage"]["candidate_count"] if graph else None}))


if __name__ == "__main__":
    main()
